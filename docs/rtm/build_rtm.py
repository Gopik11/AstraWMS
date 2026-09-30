"""Regenerate docs/rtm/AstraWMS_RTM.xlsx from the scope documents in docs/scope/*.md.

Usage: python docs/rtm/build_rtm.py [--fresh]  (requires openpyxl)

Columns derived from the documents (ID, type, area, section, text, acceptance
criteria, priority) are always refreshed. Team-maintained columns (fit-gap,
design ref, test case ID/type/phase, status, defects, owner, notes) are carried
over from the existing workbook by Req ID, including extra test-case rows.
Rows whose Req ID no longer exists in the documents move to an "Orphaned" sheet.
The previous workbook is kept as AstraWMS_RTM.bak.xlsx. --fresh ignores it.
"""
import re, glob, os, sys, shutil, argparse
from collections import Counter
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.utils import get_column_letter

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
DOCS = os.path.join(ROOT, "docs", "scope")
OUT_DIR = os.path.join(ROOT, "docs", "rtm")
OUT = os.path.join(OUT_DIR, "AstraWMS_RTM.xlsx")

AREAS = {
    "INB": "Inbound", "PUT": "Putaway & Storage", "OUT": "Outbound Order Mgmt",
    "PCK": "Picking", "SHP": "Packing, Staging & Shipping", "INV": "Inventory Management",
    "RPL": "Replenishment", "RET": "Returns", "XDK": "Cross-Docking", "VAS": "Value-Added Services",
    "KIT": "Kitting / Light Mfg", "QM": "Quality Management", "MWH": "Multi-WH & 3PL",
    "CCH": "Cold Chain & HazMat", "ADV": "Advanced Features", "INT": "ERP Integration",
    "IF": "ERP Interfaces", "NFR": "Non-Functional", "UX": "UI/UX",
    "RPT": "Reporting & Analytics",
}
LABEL_HEADERS = ("Area", "Aspect")  # tables where column 2 is a label and column 3 the requirement text
AREA_ORDER = list(AREAS.values())

def clean(s):
    s = s.replace("\\|", "|")
    s = re.sub(r"\*\*(.+?)\*\*", r"\1", s)
    s = re.sub(r"`([^`]*)`", r"\1", s)
    return s.strip()

def split_row(line):
    parts = re.split(r"(?<!\\)\|", line.strip())[1:-1]
    return [clean(p) for p in parts]

PERF_RX = re.compile(r"(p9[59]|≤\s*\d+\s*(ms|s\b|min)|readings/s|≤ 5 min\b)", re.I)

def test_profile(prefix, num, kind, text):
    if kind == "Exception":
        return "Negative / exception scenario", "SIT; UAT"
    if kind == "Interface":
        return "Interface test per ISD (positive, negative, idempotency, ordering)", "SIT"
    if prefix == "INT":
        if num >= 20:  return "Integration security (authN/authZ, encryption, audit)", "SIT"
        return "Integration (component) + end-to-end SIT", "SIT"
    if prefix == "UX":
        return "UI functional + role-based access (per module)", "SIT; UAT"
    if prefix == "RPT":
        if num >= 60:  return "Report validation (content, schedule, security)", "UAT"
        return "KPI validation (formula reconciled to source data)", "UAT"
    if prefix == "NFR":
        if num >= 120: return "API contract & governance (lint, contract tests)", "SIT"
        if num >= 100: return "Security & compliance", "SIT; UAT"
        if num <= 5:   return "Performance / scalability", "PERF"
        if num <= 30:  return "Resilience / DR", "PERF; DR"
        if num <= 47:  return "Configuration verification", "SIT"
        if num <= 66:  return "Security & audit", "SIT; UAT"
        if num == 81:  return "Accessibility (WCAG 2.2 AA)", "UAT"
        if num == 82:  return "Operability / production readiness", "PERF"
        return "Usability", "UAT"
    if PERF_RX.search(text):
        return "Performance", "PERF"
    if prefix == "ADV" and 50 <= num <= 53:
        return "Automation / MHE (emulator, FAT/SAT)", "SIT; Site"
    t = "Functional (rule test harness + E2E scenario)"
    if re.search(r"e-signature|Part 11", text, re.I):
        t += "; Regulated validation (OQ)"
    return t, "SIT; UAT"

rows, seen = [], set()
for path in sorted(glob.glob(os.path.join(DOCS, "*.md"))):
    fname = os.path.basename(path)
    section, header, prev = "", [], ""
    for line in open(path, encoding="utf-8"):
        if re.match(r"^\|[-| :]+\|\s*$", line):
            header = split_row(prev)
        prev = line
        h = re.match(r"^(#{2,3})\s+(.*)", line)
        if h:
            if re.match(r"^([A-Z]\.)?\d", clean(h.group(2))):
                section = clean(h.group(2))
            continue
        # table rows: requirements & exceptions
        m = re.match(r"^\|\s*([A-Z]{2,4})-(EX-)?(\d{2,3})\s*\|", line)
        if m:
            prefix, ex, num = m.group(1), m.group(2), int(m.group(3))
            if prefix not in AREAS: continue
            c = split_row(line)
            rid = c[0]
            if ex:
                kind = "Exception"
                desc = c[1]
                if len(c) >= 5:
                    acc = f"Detection: {c[2]} | Resolution: {c[3]} | ERP impact: {c[4]}"
                elif len(c) == 4:
                    acc = f"Detection: {c[2]} | Resolution: {c[3]}"
                else:
                    acc = f"Resolution: {c[2]}"
                pri_raw = "M"
                note = "Priority defaulted to M (exception flow; see Legend)"
            else:
                kind = "Requirement"
                desc = c[1]
                third = c[2] if len(c) > 2 else ""
                if header and header[-1] == "Pri" and len(c) > 3:
                    # multi-column table with explicit Pri column (e.g. E.4, G.2, H.3)
                    pri_raw, mid, hmid, note = c[-1], c[1:-1], header[1:-1], ""
                    if hmid[0] in LABEL_HEADERS:
                        desc = f"{mid[0]}: {mid[1]}"
                        acc = "Rule/requirement enforced as stated (positive + negative case)"
                    else:
                        desc = f"{hmid[0]}: {mid[0]}"
                        acc = " | ".join(f"{hh}: {v}" for hh, v in zip(hmid[1:], mid[1:]))
                elif re.match(r"^[MSCW]\b", third):
                    pri_raw, acc, note = third, "Rule/requirement enforced as stated (positive + negative case)", ""
                else:  # e.g. NFR-020..030: third column is the target
                    pri_raw, acc = "M", f"Target: {third}"
                    note = "Priority defaulted to M (HA/DR target; see Legend)"
            pri = pri_raw[0]
            extra = pri_raw[1:].strip(" ()")
            if extra: note = (note + "; " if note else "") + f"Priority qualifier: {extra}"
        else:
            # bullet requirements: "- INT-001 (M): text"
            b = re.match(r"^- ([A-Z]{2,4})-(\d{3}) \(([MSCW])\):\s*(.*)", line)
            i = re.match(r"^\|\s*(IF-[A-Z]+-\d{3})\s*\|", line)
            if b:
                prefix, num = b.group(1), int(b.group(2))
                rid, kind, pri, desc = f"{prefix}-{b.group(2)}", "Requirement", b.group(3), clean(b.group(4))
                acc, note = "Rule/requirement enforced as stated (positive + negative case)", ""
            elif i:
                c = split_row(line)
                prefix, num, rid, kind, pri = "IF", 0, c[0], "Interface", "M"
                desc = f"{c[2]} ({c[1]}), trigger: {c[5]}"
                acc = f"SAP: {c[3]} | Oracle: {c[4]} | Design peak: {c[6]} | Latency SLA: {c[7]}"
                note = "Priority defaulted to M (interface in catalogue D.4.2)"
            else:
                continue
        if rid in seen: continue
        seen.add(rid)
        ttype, phase = test_profile(prefix, num, kind, desc + " " + acc)
        rows.append(dict(id=rid, kind=kind, area=AREAS[prefix], section=section, doc=f"docs/scope/{fname}",
                         desc=desc, acc=acc, pri=pri, tc=f"TC-{rid}-01", ttype=ttype, phase=phase, note=note))

rows.sort(key=lambda r: (AREA_ORDER.index(r["area"]), {"Requirement": 0, "Interface": 1, "Exception": 2}[r["kind"]], r["id"]))

# ---------------- merge with existing workbook ----------------
ap = argparse.ArgumentParser(description="Regenerate the AstraWMS RTM from docs/scope/*.md")
ap.add_argument("--fresh", action="store_true", help="ignore entries in the existing workbook and start blank")
args = ap.parse_args()

cols = [  # (header, width, team-maintained input column)
    ("Req ID", 13, False), ("Type", 12, False), ("Area", 20, False), ("Section", 26, False),
    ("Source Document", 30, False), ("Requirement / Exception", 55, False),
    ("Acceptance Criteria / Expected Behaviour", 55, False), ("Priority", 9, False),
    ("Fit-Gap Disposition", 16, True), ("Design Reference", 18, True),
    ("Test Case ID", 18, False), ("Test Type", 30, False), ("Test Phase", 12, False),
    ("Test Status", 13, True), ("Defect ID(s)", 14, True), ("Owner", 16, True), ("Notes", 40, False),
]
HEADERS = [c[0] for c in cols]
PRESERVE = ["Fit-Gap Disposition", "Design Reference", "Test Case ID", "Test Type", "Test Phase",
            "Test Status", "Defect ID(s)", "Owner", "Notes"]
EXTRA_TC = "Additional Test Case"

def gen_values(r):
    return dict(zip(HEADERS, [r["id"], r["kind"], r["area"], r["section"], r["doc"], r["desc"], r["acc"], r["pri"],
                              None, None, r["tc"], r["ttype"], r["phase"], "Not Started", None, None, r["note"] or None]))

def read_sheet(wbk, name):
    if name not in wbk.sheetnames:
        return []
    sh = wbk[name]
    hdr = [c.value for c in sh[1]]
    out = []
    for vals in sh.iter_rows(min_row=2, values_only=True):
        d = dict(zip(hdr, vals))
        if str(d.get("Req ID") or "").strip():
            d["Req ID"] = str(d["Req ID"]).strip()
            out.append(d)
    return out

existing = {}
if os.path.exists(OUT) and not args.fresh:
    old = load_workbook(OUT)
    if "RTM" not in old.sheetnames:
        sys.exit(f"{OUT} has no RTM sheet; re-run with --fresh to rebuild it from scratch.")
    for d in read_sheet(old, "RTM") + read_sheet(old, "Orphaned"):
        existing.setdefault(d["Req ID"], []).append(d)

out_rows, conflicts, stats = [], [], Counter()
gen_ids = {r["id"] for r in rows}
for r in rows:
    g = gen_values(r)
    olds = sorted(existing.get(r["id"], []), key=lambda d: d.get("Type") == EXTRA_TC)  # primary row first
    if not olds:
        out_rows.append(g); stats["new"] += 1
        continue
    stats["carried over"] += 1
    for k, o in enumerate(olds):
        row = dict(g)
        if k:
            row.update({"Type": EXTRA_TC, "Test Case ID": f"TC-{r['id']}-{k + 1:02d}", "Notes": None})
            stats["additional test cases"] += 1
        for col in PRESERVE:
            if o.get(col) not in (None, ""):
                row[col] = o[col]
        if not k and o.get("Priority") not in (None, "", g["Priority"]):
            conflicts.append(f"{r['id']}: workbook priority {o['Priority']} replaced by document priority {g['Priority']}")
        out_rows.append(row)
orphans = [d for rid, lst in existing.items() if rid not in gen_ids for d in lst]
stats["orphaned"] = len(orphans)

# ---------------- workbook ----------------
F = "Arial"
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
hdr_fill = PatternFill("solid", fgColor="1F3864")
input_fill = PatternFill("solid", fgColor="FFFF00")
hdr_font = Font(name=F, bold=True, color="FFFFFF", size=10)
input_hdr_font = Font(name=F, bold=True, color="000000", size=10)
body = Font(name=F, size=9)
wrap_top = Alignment(wrap_text=True, vertical="top")

wb = Workbook()

# ---- RTM sheet
ws = wb.active
ws.title = "RTM"

def write_table(sheet, data):
    for ci, (name, width, is_input) in enumerate(cols, 1):
        cell = sheet.cell(row=1, column=ci, value=name)
        cell.font = input_hdr_font if is_input else hdr_font
        cell.fill = input_fill if is_input else hdr_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = border
        sheet.column_dimensions[get_column_letter(ci)].width = width
    sheet.row_dimensions[1].height = 32
    for ri, d in enumerate(data, 2):
        for ci, name in enumerate(HEADERS, 1):
            c = sheet.cell(row=ri, column=ci, value=d.get(name))
            c.font = body
            c.alignment = wrap_top
            c.border = border

write_table(ws, out_rows)
last = len(out_rows) + 1
ws.freeze_panes = "B2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{last}"

dv_status = DataValidation(type="list", formula1='"Not Started,In Progress,Passed,Failed,Blocked,Deferred,N/A"', allow_blank=True)
dv_fit = DataValidation(type="list", formula1='"Fit - Standard,Fit - Configuration,Gap - Extension,Gap - Process Change,Rejected,Deferred"', allow_blank=True)
dv_pri = DataValidation(type="list", formula1='"M,S,C,W"', allow_blank=False)
for dv, col in ((dv_status, "N"), (dv_fit, "I"), (dv_pri, "H")):
    ws.add_data_validation(dv)
    dv.add(f"{col}2:{col}{last}")

status_colors = {"Passed": "C6EFCE", "Failed": "FFC7CE", "Blocked": "FFC7CE", "In Progress": "FFEB9C", "Deferred": "D9D9D9", "N/A": "D9D9D9"}
for val, color in status_colors.items():
    ws.conditional_formatting.add(f"N2:N{last}", CellIsRule(operator="equal", formula=[f'"{val}"'], fill=PatternFill("solid", fgColor=color)))
ws.conditional_formatting.add(f"H2:H{last}", CellIsRule(operator="equal", formula=['"M"'], font=Font(name=F, bold=True, color="C00000")))

# ---- Orphaned sheet: rows whose Req ID is no longer in the documents (kept, never deleted)
if orphans:
    osh = wb.create_sheet("Orphaned")
    write_table(osh, orphans)
    osh.freeze_panes = "B2"

# ---- Summary sheet
sm = wb.create_sheet("Summary")
sm["A1"] = "AstraWMS: Requirements Traceability Summary"
sm["A1"].font = Font(name=F, bold=True, size=14)
sm["A2"] = "All figures are formulas over the RTM sheet and update as Test Status / Fit-Gap Disposition are maintained."
sm["A2"].font = Font(name=F, italic=True, size=9, color="595959")
heads = ["Area", "Total Items", "Requirements", "Interfaces", "Exceptions", "Must (M)", "Should (S)", "Could (C)",
         "Fit-Gap Assessed", "Test Cases", "Passed", "Failed / Blocked", "In Progress", "Not Started",
         "% Test Cases Passed", "% Fit-Gap Assessed"]
HR = 4
for ci, h in enumerate(heads, 1):
    c = sm.cell(row=HR, column=ci, value=h)
    c.font = hdr_font; c.fill = hdr_fill; c.border = border
    c.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
    sm.column_dimensions[get_column_letter(ci)].width = 26 if ci == 1 else 12
sm.row_dimensions[HR].height = 32

rng = lambda col: f"RTM!${col}$2:${col}${last}"
areas_present = [a for a in AREA_ORDER if any(r["area"] == a for r in rows)]
for i, area in enumerate(areas_present):
    r = HR + 1 + i
    a = f"$A{r}"
    primary = f'{rng("B")},"<>{EXTRA_TC}"'  # priority / fit-gap count items, not extra test-case rows
    f = {
        2: f"=C{r}+D{r}+E{r}",
        3: f'=COUNTIFS({rng("C")},{a},{rng("B")},"Requirement")',
        4: f'=COUNTIFS({rng("C")},{a},{rng("B")},"Interface")',
        5: f'=COUNTIFS({rng("C")},{a},{rng("B")},"Exception")',
        6: f'=COUNTIFS({rng("C")},{a},{rng("H")},"M",{primary})',
        7: f'=COUNTIFS({rng("C")},{a},{rng("H")},"S",{primary})',
        8: f'=COUNTIFS({rng("C")},{a},{rng("H")},"C",{primary})',
        9: f'=COUNTIFS({rng("C")},{a},{rng("I")},"?*",{primary})',
        10: f"=COUNTIFS({rng('C')},{a})",
        11: f'=COUNTIFS({rng("C")},{a},{rng("N")},"Passed")',
        12: f'=COUNTIFS({rng("C")},{a},{rng("N")},"Failed")+COUNTIFS({rng("C")},{a},{rng("N")},"Blocked")',
        13: f'=COUNTIFS({rng("C")},{a},{rng("N")},"In Progress")',
        14: f'=COUNTIFS({rng("C")},{a},{rng("N")},"Not Started")',
        15: f"=IFERROR(K{r}/J{r},0)",
        16: f"=IFERROR(I{r}/B{r},0)",
    }
    sm.cell(row=r, column=1, value=area)
    for ci, formula in f.items():
        sm.cell(row=r, column=ci, value=formula)
tr = HR + 1 + len(areas_present)
sm.cell(row=tr, column=1, value="TOTAL")
for ci in range(2, 15):
    L = get_column_letter(ci)
    sm.cell(row=tr, column=ci, value=f"=SUM({L}{HR+1}:{L}{tr-1})")
sm.cell(row=tr, column=15, value=f"=IFERROR(K{tr}/J{tr},0)")
sm.cell(row=tr, column=16, value=f"=IFERROR(I{tr}/B{tr},0)")
for r in range(HR + 1, tr + 1):
    for ci in range(1, 17):
        c = sm.cell(row=r, column=ci)
        c.font = Font(name=F, size=10, bold=(r == tr))
        c.border = border
        if ci >= 15: c.number_format = "0.0%;-0.0%;-"
        elif ci >= 2: c.number_format = "#,##0;-#,##0;-"
        if r == tr: c.fill = PatternFill("solid", fgColor="D9E1F2")
sm.cell(row=tr + 2, column=1, value="Check: RTM row count").font = Font(name=F, size=9, italic=True)
sm.cell(row=tr + 2, column=2, value=f"=COUNTA({rng('A')})").font = Font(name=F, size=9, italic=True)
sm.cell(row=tr + 2, column=3, value=f'=IF(B{tr+2}=J{tr},"OK - all rows mapped to an area","MISMATCH")').font = Font(name=F, size=9, italic=True)
sm.freeze_panes = f"B{HR+1}"

# ---- Legend sheet
lg = wb.create_sheet("Legend")
lg.column_dimensions["A"].width = 30
lg.column_dimensions["B"].width = 110
lines = [
    ("AstraWMS Requirements Traceability Matrix", None, "title"),
    ("Source", "AstraWMS Scope & Solution Definition v1.0 (docs/scope/*.md), baseline 2026-09-30. Every requirement ID, exception code and interface ID in the document set is one row on the RTM sheet.", None),
    ("How to use", "Maintain the yellow-header columns on the RTM sheet (Fit-Gap Disposition, Design Reference, Test Status, Defect ID(s), Owner); Test Case ID, Test Type, Test Phase and Notes may also be edited. Summary recalculates automatically. Add further test cases for a requirement as extra rows with the same Req ID; they are typed 'Additional Test Case' on the next regeneration.", None),
    ("Regenerating", "python docs/rtm/build_rtm.py refreshes the document-derived columns (Req ID to Priority) and keeps everything the team entered, matched by Req ID. Priority always follows the scope documents: change it there. Rows whose Req ID was removed from the documents move to the Orphaned sheet. The previous file is saved as AstraWMS_RTM.bak.xlsx. --fresh starts from a blank matrix.", None),
    ("", None, None),
    ("Column", "Meaning", "head"),
    ("Req ID", "Identifier from the scope documents: <AREA>-<NNN> requirement, <AREA>-EX-<NN> exception flow, IF-<FLOW>-<NNN> ERP interface."),
    ("Type", "Requirement | Exception | Interface | Additional Test Case (extra test-case row for a Req ID; excluded from item, priority and fit-gap counts)."),
    ("Area / Section / Source Document", "Functional area (from the ID prefix), the heading it appears under, and the Markdown file that defines it."),
    ("Requirement / Exception", "Requirement text, exception name, or interface object/direction/trigger as written in the scope."),
    ("Acceptance Criteria", "Rules: enforced as stated (positive + negative case). Exceptions: detection, resolution and ERP impact from the exception table. Interfaces: SAP/Oracle technology, design peak volume and latency SLA. HA/DR NFRs: the stated target."),
    ("Priority", "MoSCoW: M Must, S Should, C Could, W Won't (document §0.4)."),
    ("Fit-Gap Disposition (input)", "Blueprint outcome (§I.1.2): Fit - Standard, Fit - Configuration, Gap - Extension, Gap - Process Change, Rejected, Deferred."),
    ("Design Reference (input)", "Process Design Document (PDD) or Interface Specification Document (ISD) reference."),
    ("Test Case ID", "Generated placeholder TC-<ReqID>-01; replace/extend with IDs from the test management tool."),
    ("Test Type / Test Phase", "Proposed from the Testing Strategy (§I.2): functional, exception, integration, interface, performance, resilience/DR, security & audit, automation, regulated validation."),
    ("Test Status (input)", "Not Started, In Progress, Passed, Failed, Blocked, Deferred, N/A."),
    ("", None, None),
    ("Example (filled row)", "Req ID INB-002 · Fit-Gap Disposition: Fit - Configuration · Design Reference: PDD-INB-01 §3.2 · Test Status: Passed · Defect ID(s): DEF-0142 (closed) · Owner: Receiving Process Owner", None),
    ("", None, None),
    ("Assumptions", None, "head"),
    ("A1", "Exception codes carry no priority in the scope; they are set to M because §I.2 requires every exception code to be covered by a test case."),
    ("A2", "NFR-020 to NFR-030 (HA/DR) list a target instead of a priority; they are set to M."),
    ("A3", "Interfaces IF-* are set to M; drop to S/C per site if an interface is not used (e.g. kitting IF-KIT-001 at sites without kitting)."),
    ("A4", "Where a priority carries a qualifier (e.g. 'M (if sorter in scope)'), the letter is used and the qualifier is kept in Notes."),
    ("A5", "Test Type and Phase are proposals derived from the requirement wording (latency/throughput targets → Performance; e-signature/Part 11 → Regulated validation). Confirm during test planning."),
    ("A6", "KPI and report rows (RPT-*) take the KPI formula, target and grain, or the report frequency, audience and contents, as acceptance criteria. Web UI module rows (UX-*) take the key screens and primary roles."),
    ("Coverage", "Traced: every numbered requirement, exception code and interface, plus §D.5 Integration Security (INT-020–027), §E.4 Security & Compliance (NFR-100–109), §E.6 API Governance (NFR-120–129), §G.2 Web UI Modules (UX-001–017), §H.3–H.6 KPIs and reports (RPT-001–070). Not traced as separate rows: descriptive design tables (§F architecture, §D.1–D.4 mappings, §G.3–G.6 RF flows/voice/RBAC/dashboards, §H.7 data warehouse), which are verified through the requirements that reference them."),
]
r = 1
for item in lines:
    k, v = item[0], item[1]
    style = item[2] if len(item) > 2 else None
    ca = lg.cell(row=r, column=1, value=k or None)
    cb = lg.cell(row=r, column=2, value=v)
    if style == "title":
        ca.font = Font(name=F, bold=True, size=14)
    elif style == "head":
        for c in (ca, cb):
            c.font = hdr_font; c.fill = hdr_fill
    else:
        ca.font = Font(name=F, bold=True, size=10)
        cb.font = Font(name=F, size=10)
    cb.alignment = Alignment(wrap_text=True, vertical="top")
    ca.alignment = Alignment(vertical="top")
    r += 1

for pos, name in enumerate(["Legend", "Summary", "RTM"]):  # Orphaned (if any) ends up last
    wb.move_sheet(name, offset=pos - wb.sheetnames.index(name))
wb.active = 0
os.makedirs(OUT_DIR, exist_ok=True)
from openpyxl.workbook.properties import CalcProperties
wb.calculation = CalcProperties(fullCalcOnLoad=True)
tmp = OUT + ".tmp"
wb.save(tmp)
try:
    if os.path.exists(OUT):
        shutil.copy2(OUT, os.path.join(OUT_DIR, "AstraWMS_RTM.bak.xlsx"))
    os.replace(tmp, OUT)
except PermissionError:
    os.remove(tmp)
    sys.exit(f"Cannot write {OUT}: close it in Excel and re-run. Nothing was changed.")

print(f"{len(rows)} items ({dict(Counter(r['kind'] for r in rows))}), {len(out_rows)} RTM rows")
print("Merge:", ", ".join(f"{k} {v}" for k, v in stats.items()) if existing else "fresh build (no existing entries used)")
for c in conflicts:
    print("  priority:", c)
for d in orphans:
    print(f"  orphaned: {d['Req ID']} -> Orphaned sheet")
